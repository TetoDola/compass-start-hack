from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("/Users/benjaminauer/projects/unriskomega-case-document/UNRISKOMEGA_Case_Knowledge_Base.docx")

NAVY = "0F1B2D"
ORANGE = "F47A20"
PALE_BLUE = "EEF3F8"
PALE_ORANGE = "FFF3E9"
LIGHT_GRAY = "F5F6F7"
BORDER = "D9D9D9"
MID_GRAY = "666666"
BLACK = "000000"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def set_col_widths(table, widths) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)


def add_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_run_font(run, name="Aptos", size=None, bold=None, color=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def add_para(doc, text="", *, style=None, bold_lead=None, italic=False, space_after=6, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, italic=italic)
    else:
        run = p.add_run(text)
        set_run_font(run, italic=italic)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.25 + level * 0.18)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_number(doc, text, level=0):
    p = doc.add_paragraph(style="List Number" if level == 0 else "List Number 2")
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_number_literal(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    run = p.add_run(f"{number}.\t{text}")
    set_run_font(run)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_table(doc, headers, rows, widths=None, header_fill=NAVY, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, header in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, header_fill)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(header))
        set_run_font(run, size=9, bold=True, color=WHITE)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        fill = PALE_BLUE if r_idx % 2 else WHITE
        for i, value in enumerate(row):
            cell = cells[i]
            set_cell_shading(cell, fill)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            set_run_font(run, size=font_size)
    if widths:
        set_col_widths(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(3)
    return table


def add_label(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text.upper())
    set_run_font(run, size=8.5, bold=True, color=ORANGE)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.different_first_page_header_footer = True

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.7)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(6)

title_style = styles["Title"]
title_style.font.name = "Aptos Display"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
title_style.font.size = Pt(30)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor.from_string(BLACK)
title_style.paragraph_format.space_after = Pt(12)
title_style_ppr = title_style._element.get_or_add_pPr()
title_style_borders = title_style_ppr.find(qn("w:pBdr"))
if title_style_borders is not None:
    title_style_ppr.remove(title_style_borders)

for level, size, before, after in ((1, 18, 16, 7), (2, 13, 12, 5), (3, 11, 8, 3)):
    style = styles[f"Heading {level}"]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(BLACK)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

for list_style in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
    styles[list_style].font.name = "Aptos"
    styles[list_style]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles[list_style]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles[list_style].font.size = Pt(10.5)

# Header and footer
header = section.header
hp = header.paragraphs[0]
hp.text = "UNRISKOMEGA Case Knowledge Base"
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
for run in hp.runs:
    set_run_font(run, size=8.5, color=MID_GRAY)

footer = section.footer
fp = footer.paragraphs[0]
fp.add_run("Prepared 18 September 2026   |   ")
for run in fp.runs:
    set_run_font(run, size=8.5, color=MID_GRAY)

# Keep the cover visually quiet while preserving the running header and page
# numbers on the report pages that follow.
section.first_page_header.paragraphs[0].text = ""
section.first_page_footer.paragraphs[0].text = ""
add_page_field(fp)
for run in fp.runs:
    set_run_font(run, size=8.5, color=MID_GRAY)

# Cover page
spacer = doc.add_paragraph()
spacer.paragraph_format.space_after = Pt(70)
add_label(doc, "Case report and implementation guide")
title = doc.add_paragraph(style="Title")
title.add_run("UNRISKOMEGA Case Knowledge Base")
remove_paragraph_borders(title)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
run = subtitle.add_run("From Ping to Pitch in 60 Seconds")
set_run_font(run, name="Aptos Display", size=17, bold=False, color=NAVY)
meta = doc.add_paragraph()
meta.paragraph_format.space_after = Pt(18)
run = meta.add_run("Consolidated from the official case deck, three recorded conversations, the challenge repository audit, and prior case discussions")
set_run_font(run, size=11.5, color=MID_GRAY)

add_para(doc, "Purpose", bold_lead="Purpose", space_after=4)
add_para(
    doc,
    "This report gives the team one reliable source for the UNRISKOMEGA challenge. It separates verified facts from transcript-derived clarifications and from recommended implementation choices. The core conclusion is that the strongest solution is a one-click, client-specific briefing backed by deterministic calculations and inspectable evidence.",
    space_after=8,
)
add_para(doc, "Prepared for hackathon planning and prototype delivery", italic=True, space_after=0)
doc.add_page_break()

# Contents overview
add_heading(doc, "Document guide", 1)
add_para(doc, "The report is organized around the decisions the team must make, from case understanding through delivery and judging.")
add_table(
    doc,
    ["Section", "What it establishes"],
    [
        ("Executive summary", "The recommended solution and the reasons it fits the case"),
        ("Company and product context", "UNRISKOMEGA, URO Advisor Pro and the target advisory workflow"),
        ("Sponsor requirements", "Required output, clarified scope and judging criteria"),
        ("Data and repository", "Available data, measured coverage and material quality risks"),
        ("Product and architecture", "Briefing design, grounding model, APIs and deployment constraints"),
        ("Demo and delivery", "CASE-012 story, acceptance tests, build order and pitch strategy"),
        ("Appendices", "Open decisions, source provenance and transcript-specific contributions"),
    ],
    widths=[1.9, 4.8],
)

add_heading(doc, "Executive summary", 1)
add_para(
    doc,
    "UNRISKOMEGA wants an AI briefing assistant inside URO Advisor Pro that prepares an investment adviser for a client conversation in seconds. The adviser should trigger it from the existing workflow, read it in about one minute, and understand recent portfolio developments, current portfolio health, relevant outlook, and the next actions worth discussing.",
)
add_para(
    doc,
    "The most credible implementation uses deterministic code for portfolio joins, calculations, thresholds, rule checks, freshness and ranking. A language model should receive only a compact evidence ledger and turn approved evidence into concise prose. Every sentence should expose its source, date and confidence. When a source is absent or stale, the interface should say so instead of filling the gap.",
)
add_para(
    doc,
    "The recommended demo centers on CASE-012. A client note records an upcoming need for about CHF 15,000 of liquidity, while the portfolio is concentrated in USD equities and has 21 active violations. This connects portfolio facts, client context and actionable discussion points in one clear story. A second unseen client should prove that the system is not hardcoded.",
)
add_table(
    doc,
    ["Decision", "Recommendation", "Reason"],
    [
        ("Primary product", "One-click 60-second briefing", "Directly matches the required workflow and 85 percent core judging scope"),
        ("AI role", "Narrative composition from approved evidence", "Keeps calculations auditable and limits unsupported claims"),
        ("Primary demo", "CASE-012", "Combines a cash need, concentration, risk and violations in one client-specific story"),
        ("External sources", "One dated news item and one approved house-view sample", "Proves the integration pattern without relying on a fragile live feed"),
        ("First bonus", "Ex-custody PDF import", "Extends the same portfolio model and solves a confirmed manual-work problem"),
        ("Voice copilot", "Optional extension after the button flow works", "Recording and privacy constraints make it unsuitable as the only trigger"),
    ],
    widths=[1.25, 2.25, 3.2],
)

add_heading(doc, "Company and product context", 1)
add_heading(doc, "UNRISKOMEGA", 2)
add_para(
    doc,
    "UNRISKOMEGA is a Swiss WealthTech company founded in 2016 and organized as a joint venture of four owner-operated companies. The official deck lists headquarters in Kloten and development locations in Perg and Timisoara. It reports 18 employees, more than 25 years of digital investment-advisory expertise, more than 25 bank clients, about 1,000 active advisers, more than 500,000 portfolios and more than 100,000 instruments across relevant asset classes.",
)
add_para(
    doc,
    "The presentation speaker described roughly 25 people overall and explained that 18 represented the full-time-equivalent product team. He said the first client went live about eight years earlier. These statements explain the apparent difference between headcount figures but remain transcript-derived rather than formal company disclosures.",
)
add_table(
    doc,
    ["Company fact", "Case material"],
    [
        ("Business", "B2B software for banks and wealth managers"),
        ("Core expertise", "Digital investment advisory with FIDLEG and MiFID support"),
        ("Locations", "Kloten, Perg and Timisoara"),
        ("Clients", "More than 25 banks across universal, cantonal, regional and private banking"),
        ("Platform reach", "About 1,000 advisers, 500,000-plus portfolios and 100,000-plus instruments"),
    ],
    widths=[1.8, 4.9],
)

add_heading(doc, "URO Advisor Pro", 2)
add_para(
    doc,
    "URO Advisor Pro supports the end-to-end investment-advisory process. Its functions include investor profiling, simulation and documentation, order preparation, portfolio review, portfolio monitoring, a risk engine, and a rule engine for contractual and regulatory checks. It integrates with the Finnova core-banking system and can connect bank-specific systems and additional data sources such as ESG data, fund breakdowns, product documents and digital signatures.",
)
add_para(
    doc,
    "The case targets the adviser-facing B2B product rather than the self-service products. The relevant client segment starts around affluent clients with several hundred thousand Swiss francs invested and extends into the low millions. The sponsor advised against optimizing for ultra-high-net-worth clients, whose service model is more personal and whose advisers manage fewer relationships.",
)

add_heading(doc, "The 60 second problem", 1)
add_para(
    doc,
    "Investment advisers manage hundreds of relationships. An unannounced call can arrive after a market move, while portfolio analytics, CRM notes, proposals, news and the bank's house view remain spread across separate systems. The adviser must turn those sources into a useful explanation before the conversation starts. Manual preparation is too slow, and the quality of the response depends heavily on the individual adviser.",
)
add_table(
    doc,
    ["Current condition", "Operational effect", "Required product response"],
    [
        ("Calls or walk-ins arrive without warning", "One or two minutes of preparation time", "Generate a briefing from a clear entry point inside URO"),
        ("Data is distributed", "The adviser searches several applications and documents", "Combine portfolio, CRM, market and house-view evidence"),
        ("Storytelling is manual", "Important drivers or opportunities may be missed", "Rank the few facts that matter for this client"),
        ("Time window closes", "Client confidence and adviser reputation suffer", "Return a readable answer within seconds"),
    ],
    widths=[2.0, 2.3, 2.4],
)
add_para(
    doc,
    "The case deck frames the business value as saved preparation time, higher adviser confidence, fewer surprises, fewer missed reinvestment opportunities and more consistent conversation quality across the advisory team. The transcript adds a reputation benefit: the client experiences an adviser who appears prepared and informed rather than someone searching through several tools during the call.",
)

add_heading(doc, "Sponsor requirements and clarifications", 1)
add_heading(doc, "Required briefing", 2)
add_para(doc, "The minimum viable briefing must answer four questions in a client-specific way:")
add_number(doc, "What happened recently, including portfolio value or performance measures and the main drivers.")
add_number(doc, "What matters now, including portfolio health, concentration, strategic allocation deviations and active risk or suitability violations.")
add_number(doc, "What could happen next, using only market news and CIO or house-view evidence that maps to actual holdings or exposures.")
add_number(doc, "What the adviser should discuss next, phrased as options or preparation steps rather than automated investment instructions.")
add_para(
    doc,
    "The sponsor suggested a short bullet format because advisers are already overloaded. Three memorable points per major section are preferable to a dense report. The deck specifies roughly 60 seconds of reading time. In Q&A, the sponsor described generation in about 10 seconds as an impressive target.",
)

add_heading(doc, "Workflow scope", 2)
add_bullet(doc, "The core flow starts with a button or comparable explicit trigger in URO Advisor Pro.")
add_bullet(doc, "The workflow must support incoming client calls, adviser-initiated calls, scheduled remote meetings and physical meetings.")
add_bullet(doc, "The expected operating setting is an adviser at a workstation with access to the client and portfolio view.")
add_bullet(doc, "Live call recognition and coaching are valid creative extensions, but they do not replace the button-triggered core case.")
add_bullet(doc, "Some banks restrict or prohibit call recording. A voice feature therefore needs a bank-controlled opt-in and a non-recording fallback.")
add_bullet(doc, "Text must remain available because some users will find graphs or conversational interfaces distracting. Banks may want to disable optional views.")

add_heading(doc, "Information sources", 2)
add_para(
    doc,
    "Production inputs would include URO portfolio and position data, CRM notes, proposals, bank-approved market data, paid news providers, internal research, and CIO or house-view material. Sponsor Q&A described house views as decentralized PDFs, spreadsheets, intranet pages and public investment outlooks. For the prototype, the team may use free public news and a dated public or mocked house-view document, provided the source and date remain visible.",
)
add_para(
    doc,
    "The adviser must be able to trust the briefing because the bank chooses which sources are approved. Credibility therefore depends on provenance, dates, explicit uncertainty and a visible failure state when information cannot be verified.",
)

add_heading(doc, "Judging and delivery expectations", 2)
add_table(
    doc,
    ["Criterion", "Weight", "Evidence expected on stage"],
    [
        ("Problem fit and business value", "25%", "A useful one-click briefing that addresses development, health, outlook and action"),
        ("Implementation quality and robustness", "25%", "Stable end-to-end data flow, sensible errors, latency and modular integrations"),
        ("AI relevance and grounding", "20%", "Concise client-specific claims and practical actions with no generic or invented content"),
        ("User experience", "15%", "Readable in about one minute and placed naturally in the adviser workflow"),
        ("Bonus features", "15%", "Ex-custody import, grounded follow-up chat or a useful alternative presentation"),
    ],
    widths=[2.05, 0.65, 4.0],
)
add_para(
    doc,
    "The sponsor said a working prototype matters more than slideware and that a surprise test client would be supplied shortly before the presentation. The core workflow accounts for 85 percent of the score. Bonus work should begin only after the briefing works repeatedly on known and unseen data.",
)

add_heading(doc, "Repository and available data", 1)
add_para(
    doc,
    "The reviewed challenge repository is a data and design pack rather than a starter application. The audit covered commit 134d6ff9343b8703a63a7279043c57c09345d45b. It contains client and reference JSON, field documentation, three URO Advisor Pro screenshots, brand assets and ten synthetic German quarterly portfolio reports for the ex-custody bonus. It does not contain application code, market news, a house view, API contracts, generated briefing examples or automated tests.",
)
add_table(
    doc,
    ["Core collection", "Count"],
    [
        ("Clients", "47"),
        ("Portfolios", "57"),
        ("Security positions", "703"),
        ("Cash or account positions", "123"),
        ("Proposals", "206"),
        ("Transactions", "1,274"),
        ("Active suitability violations", "180"),
        ("Client notes", "153"),
        ("Client tags", "73"),
    ],
    widths=[4.9, 1.8],
)
add_table(
    doc,
    ["Reference collection", "Count"],
    [
        ("Securities", "504"),
        ("Fund look-through mappings", "48,101"),
        ("Suitability rules", "54"),
        ("Risk profiles", "5"),
        ("ESG profiles", "2"),
        ("Investment services", "7"),
        ("Strategies", "6"),
        ("Strategic asset allocations", "16"),
        ("Recommendation-list securities", "232"),
    ],
    widths=[4.9, 1.8],
)
add_para(
    doc,
    "Twenty-six clients have at least one active violation, 29 have proposals, 28 have transactions, all 47 have notes and 35 have tags. All reporting currencies are CHF. The dataset includes four company clients. Client-level assets under management total about CHF 31.77 million, with a median of about CHF 479,595.",
)

add_heading(doc, "Material data constraints", 1)
add_table(
    doc,
    ["Constraint", "Observed evidence", "Required control"],
    [
        ("Null handling", "206 explicit null values across 15 paths despite documentation saying missing values are omitted", "Treat arrays and scalars as nullable and validate before access"),
        ("Missing performance field", "No portfolio has PerformanceYTD", "Describe monthly values as value changes unless cash-flow-adjusted returns exist"),
        ("Missing task source", "No open-task collection exists", "Display No task source connected rather than infer tasks"),
        ("Stale prices", "59 held securities are older than 30 days and 56 are older than one year", "Attach dates and block exact trade suggestions when freshness fails"),
        ("Duplicate identifiers", "Eleven ISINs occur twice across currency-specific rows", "Use SecurityId as the primary join key"),
        ("Aggregation risk", "A client has consolidated and ordinary portfolios", "Do not add consolidated views to underlying accounts without hierarchy metadata"),
        ("Unit inconsistency", "Position weights use 0 to 1 while fund look-through uses 0 to 100", "Normalize units explicitly and test them"),
        ("Mixed taxonomy", "Codes, descriptions and statuses use several languages", "Normalize for logic and retain original wording for evidence"),
    ],
    widths=[1.25, 3.05, 2.4],
    font_size=8.6,
)
add_para(
    doc,
    "These constraints make data quality part of the product, not an implementation detail. Every evidence record should carry a source, an as-of date, a measured value, a comparison threshold where relevant, and a freshness or confidence status.",
)

add_heading(doc, "Recommended product experience", 1)
add_heading(doc, "Briefing structure", 2)
add_para(
    doc,
    "Place a Generate briefing action in the client and portfolio header. Return one card that fits within a normal laptop viewport and contains about 130 to 170 words. The card should make facts, interpretation and suggested actions visually distinct without forcing the adviser into a separate application.",
)
add_table(
    doc,
    ["Briefing section", "Content rule", "Evidence rule"],
    [
        ("What happened", "One headline and no more than three drivers", "State the period, measure and data date"),
        ("What matters now", "No more than three prioritized risks or client circumstances", "Show values, limits and linked notes or positions"),
        ("What could happen next", "Only context tied to actual holdings or exposures", "Show the dated article or approved house view"),
        ("What to discuss", "Two or three adviser options", "State assumptions and avoid automatic execution language"),
    ],
    widths=[1.4, 2.75, 2.55],
)
add_para(
    doc,
    "An evidence drawer should reveal the source system, as-of date, value, threshold, freshness, confidence and linked position or note. The normal view stays concise; the drawer gives the adviser and judge a way to inspect the reasoning.",
)

add_heading(doc, "Insight prioritization", 2)
add_para(doc, "Score each candidate insight using a transparent model:")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
run = p.add_run("priority = urgency × materiality × client relevance × actionability × confidence")
set_run_font(run, name="Aptos Display", size=12, bold=True, color=NAVY)
add_bullet(doc, "Urgency covers active violations, near-term cash needs, maturities and pending proposals.")
add_bullet(doc, "Materiality measures the exposure or value relative to the portfolio or client assets.")
add_bullet(doc, "Client relevance uses notes, preferences, risk profile, ESG profile and stated objectives.")
add_bullet(doc, "Actionability asks whether the adviser has a concrete next step.")
add_bullet(doc, "Confidence falls when data is stale, ambiguous, missing or internally inconsistent.")
add_para(doc, "Freshness, missing thresholds and ambiguous portfolio scope are hard gates before ranking. A high score must never override an unsafe input.")

add_heading(doc, "Technical architecture", 1)
add_table(
    doc,
    ["Layer", "Responsibility", "Prototype requirement"],
    [
        ("Ingestion", "Read client JSON, reference data, external context and optional PDF statements", "Schema validation, null safety and repeatable unseen-file support"),
        ("Canonical model", "Normalize clients, portfolios, positions, notes, proposals and evidence", "Stable IDs, unit conversion and portfolio-scope rules"),
        ("Analytics", "Calculate concentration, changes, deviations, violations, liquidity and candidate actions", "Deterministic code with dated inputs"),
        ("Relevance ranking", "Select the few facts that matter for this client and moment", "Transparent priority factors and hard safety gates"),
        ("Evidence ledger", "Record claim, value, source, date, threshold and confidence", "One or more evidence records for every sentence"),
        ("Narrative composer", "Express approved evidence in the four-part briefing schema", "Structured output and deterministic fallback"),
        ("Validator", "Reject unsupported numbers or claims", "Sentence-to-evidence mapping and word-budget check"),
        ("URO interface", "Present the briefing and drill-down evidence", "One click, one viewport and clear unavailable states"),
    ],
    widths=[1.2, 3.0, 2.5],
    font_size=8.7,
)
add_heading(doc, "Grounding and compliance controls", 2)
add_bullet(doc, "Keep joins, arithmetic, thresholds and freshness checks outside the language model.")
add_bullet(doc, "Give the model only ranked evidence and require a fixed output schema.")
add_bullet(doc, "Reject or hide any sentence that cannot point to evidence.")
add_bullet(doc, "Use a deterministic template when model output fails validation or the model is unavailable.")
add_bullet(doc, "Label suggested actions as adviser discussion prompts, not automatic investment advice.")
add_bullet(doc, "Mask realistic account and IBAN-like identifiers before logging or model use.")
add_bullet(doc, "Make missing news, house view or task data explicit rather than implied.")

add_heading(doc, "Hosting and integration constraints", 1)
add_para(
    doc,
    "Sponsor Q&A clarified that UNRISKOMEGA ships software that banks or their application-management providers deploy. UNRISKOMEGA does not operate the product as a shared software service and does not need to own client data. Production design should therefore assume deployment within bank-controlled infrastructure and access through bank-approved interfaces.",
)
add_para(
    doc,
    "Swiss hosting and organizational control may matter for some banks. The conversation mentioned private Microsoft Copilot environments, institutions where public assistants are disabled, and the possible use of locally hosted models. These comments describe practical concerns rather than a single binding rule for every bank. The architecture should remain model-agnostic and support private deployment, regional data residency, permission-aware retrieval, limited retention and auditable source access.",
)
add_table(
    doc,
    ["Concern", "Design implication"],
    [
        ("Client data remains with the bank", "Use bank-side deployment and do not require a central UNRISKOMEGA client-data store"),
        ("Approved sources differ by bank", "Configure source adapters and permissions per institution"),
        ("Voice recording may be restricted", "Keep the explicit button flow and make live listening optional"),
        ("Model choice may be constrained", "Support private or local models behind the same structured contract"),
        ("Advisory output is regulated", "Preserve the human review step and record the evidence used"),
    ],
    widths=[2.35, 4.35],
)

add_heading(doc, "Real time assistant extension", 1)
add_para(
    doc,
    "A related discussion identified the broader product category as real-time agent assist, knowledge assist or a live call copilot. The common pattern is live transcription, question or intent detection, permission-aware retrieval from internal sources, and a private, cited answer shown to the employee. Applied here, an incoming caller could be matched to a client record and the same briefing engine could prepare context before the adviser answers.",
)
add_para(
    doc,
    "This is a credible future direction, but the sponsor classified it as a creative step beyond the core case. The immediate product should expose a stable briefing API that both the button flow and a future voice trigger can call. This avoids coupling the analytical engine to call recording and lets banks activate the extension only where consent, privacy and telephony integration permit it.",
)
add_heading(doc, "Requirements for a future live call copilot", 2)
add_bullet(doc, "Detect customer questions automatically but allow the adviser to trigger or refresh manually.")
add_bullet(doc, "Return short, speakable answers with visible citations and permission-aware retrieval.")
add_bullet(doc, "Show an explicit unable to verify state when evidence is insufficient.")
add_bullet(doc, "Meet low-latency targets without skipping provenance or freshness checks.")
add_bullet(doc, "Support meeting consent, recording controls, retention limits and bank-specific data residency.")

add_heading(doc, "Primary demo case", 1)
add_heading(doc, "CASE 012", 2)
add_para(
    doc,
    "CASE-012 gives the prototype a coherent conversation rather than a generic market summary. The client has an upcoming liquidity need and a concentrated portfolio with multiple active rule violations. The briefing can connect those facts to specific preparation steps without pretending that the repository contains complete market or house-view data.",
)
add_table(
    doc,
    ["Briefing question", "CASE-012 evidence", "Safe interpretation"],
    [
        ("What happened", "Portfolio value moved from CHF 44,794 in June to CHF 44,396 in July", "About a 0.9 percent value decline; do not call it cash-flow-adjusted performance"),
        ("What matters", "Volatility 18.4 percent, VaR 20.8 percent and 21 active violations", "Risk, foreign-currency exposure and single-instrument concentration require review"),
        ("Client context", "A March note records a need for about CHF 15,000 for a Q1 tax payment", "The liquidity need is roughly one third of current assets and should drive the conversation"),
        ("Concentration", "Alphabet 25.5 percent, global equity ETF 23.8 percent, Apple 17.8 percent, Swiss dividend ETF 17.1 percent and Amazon 15.1 percent", "The portfolio depends heavily on a small set of equity positions and USD exposure"),
        ("What to discuss", "Confirm timing and amount, verify liquidity, review de-risking and refresh prices", "Prepare options for adviser review; do not generate exact trades from stale data"),
    ],
    widths=[1.25, 3.15, 2.3],
    font_size=8.7,
)
add_para(
    doc,
    "The forward-looking section needs an approved, dated house-view sample and one holding-linked news event. The UI should label both as prototype sources. The repository alone cannot support a current outlook.",
)
add_heading(doc, "Additional stress cases", 2)
add_bullet(doc, "CASE-038 tests two portfolios, one consolidated view, 18 violations, 17 proposals and several ESG or preference notes.")
add_bullet(doc, "CASE-028 tests a 95.9 percent single-equity concentration despite no active violation, proving the assistant does more than repeat the alert feed.")

add_heading(doc, "Acceptance tests", 1)
add_table(
    doc,
    ["Test", "Pass condition"],
    [
        ("Unseen client", "A new file with the supplied shape runs without code changes"),
        ("Null collections", "Missing or null arrays do not crash ingestion"),
        ("Security join", "Currency-specific duplicates resolve by SecurityId rather than ISIN alone"),
        ("Freshness", "Stale prices are flagged and block exact trade quantities"),
        ("Portfolio scope", "Consolidated and underlying portfolios are not double counted"),
        ("Evidence coverage", "Every briefing sentence has at least one evidence record"),
        ("Missing source", "The interface states when news, house view or tasks are unavailable"),
        ("Reading budget", "The briefing remains within about 130 to 170 words"),
        ("PDF reconciliation", "Imported positions reconcile to reported totals or stop for human review"),
        ("Fallback", "A deterministic template is returned when the model fails"),
    ],
    widths=[1.65, 5.05],
)

add_heading(doc, "Bonus scope", 1)
add_heading(doc, "Ex custody PDF import", 2)
add_para(
    doc,
    "The ten supplied reports are structured, text-native German PDFs with a consistent eight-page layout. They cover performance, allocation, ESG, detailed cash and security positions, and transactions. Their reported assets range from CHF 948,090 to CHF 9,203,755, and 2025 time-weighted returns range from 2.41 to 7.83 percent.",
)
for step_no, step_text in enumerate(
    [
        "Extract text and tables, using OCR only when no usable text layer exists.",
        "Parse Swiss number formats, currencies, percentages, ISINs and account identifiers.",
        "Reconcile total wealth and position weights before creating a portfolio.",
        "Match instruments using ISIN plus currency and send ambiguous matches to review.",
        "Mask IBANs and account identifiers before logging or model use.",
        "Create a labelled virtual ex-custody portfolio only after adviser confirmation.",
    ],
    start=1,
):
    add_number_literal(doc, step_no, step_text)
add_para(
    doc,
    "This bonus addresses a confirmed business problem. Clients often hold assets with several banks, there is no universal open API standard, and advisers perform significant manual work to create a complete view. The sponsor said an end-to-end button-triggered workflow would be presentable to customers and could become the starting point for deeper integration.",
)
add_heading(doc, "Other extensions", 2)
add_bullet(doc, "Grounded follow-up chat that queries the existing evidence ledger rather than the raw portfolio.")
add_bullet(doc, "Optional visual or graph view behind a toggle, with text remaining the default.")
add_bullet(doc, "Adviser feedback such as Relevant, Incorrect or Missing context to support pilot learning.")
add_bullet(doc, "Live call copilot using the same briefing API where bank policy permits it.")

add_heading(doc, "Build and demo plan", 1)
add_heading(doc, "Twelve hour implementation order", 2)
add_table(
    doc,
    ["Time", "Outcome"],
    [
        ("Hour 0 to 1", "Confirm sponsor workflow, assign owners, freeze the value claim and write the demo route"),
        ("Hour 1 to 2", "Parse the data and produce a terminal evidence bundle for CASE-012"),
        ("Hour 2 to 4", "Complete ingestion, analytics, evidence ledger and the four-part briefing"),
        ("Hour 4 to 6", "Build the URO-style UI, evidence drawer and cached external context"),
        ("Hour 6 to 8", "Add unseen-file handling, missing-data states and repeatable tests"),
        ("Hour 8 to 9", "Decide whether the core is stable enough to begin PDF import"),
        ("Hour 9 to 10", "Freeze data, cache sources, verify numbers and prepare fallback media"),
        ("Hour 10 to 11", "Finish the short deck, submission text and Q&A answers"),
        ("Hour 11 to 12", "Run full rehearsals from the presentation machine and stop adding features"),
    ],
    widths=[1.25, 5.45],
)
add_para(
    doc,
    "At hour four, a teammate who did not build the system should be able to open CASE-012, click once and explain the result. At hour eight, scope should freeze. Bonus work should start only after the primary flow has passed five consecutive runs and unseen input works.",
)

add_heading(doc, "Demo route and pitch", 2)
for step_no, step_text in enumerate(
    [
        "Open with an unexpected client call and the adviser having no time to prepare.",
        "Show the URO client or portfolio view and trigger Generate briefing.",
        "Read the CASE-012 story: value change, cash need, concentration, violations and next discussion steps.",
        "Open one evidence record to show source, value, date, threshold and freshness.",
        "Load a second unseen client and show the same pipeline and honest missing-source states.",
        "Explain which elements are real, cached, mocked and planned for production.",
    ],
    start=1,
):
    add_number_literal(doc, step_no, step_text)
add_para(
    doc,
    "The pitch should communicate that deterministic calculations produce the facts and the model writes from an evidence ledger. The jury's practical question is whether a bank could pilot the feature soon and trust an adviser to use it before a real client conversation.",
)

add_heading(doc, "Pilot proposal", 2)
add_para(
    doc,
    "A sensible next step is a four-week pilot with five to ten advisers. Measure briefing-generation time, perceived usefulness, edit rate, evidence opens, missing-context flags, unsupported-claim rate and qualitative confidence before calls. Connect only approved sources and keep human review mandatory throughout the pilot.",
)

add_heading(doc, "Open decisions", 1)
add_table(
    doc,
    ["Decision", "Why it remains open", "Recommended default"],
    [
        ("Definition of current data", "Price, news and house-view freshness may follow different rules", "Configure per source and show the date everywhere"),
        ("Portfolio hierarchy", "The dataset does not identify consolidation relationships fully", "Treat consolidated portfolios as non-additive views"),
        ("Approved next best action boundary", "Banks differ in policy and advisory controls", "Use discussion prompts and require adviser review"),
        ("External data provider", "The prototype pack has no licensed feed", "Use one cached dated source and one mock house-view adapter"),
        ("Model and hosting", "Bank policy varies", "Keep a model-agnostic interface and support private deployment"),
        ("Voice activation", "Recording and consent rules differ", "Keep voice optional and preserve manual trigger"),
        ("PDF generality", "Ten reports share one template", "Claim support for the supplied layout, not arbitrary statements"),
    ],
    widths=[1.55, 2.85, 2.3],
    font_size=8.8,
)

add_heading(doc, "Source provenance", 1)
add_para(
    doc,
    "Facts in this report were reconciled across the official deck, three Parakeet transcripts and prior repository analysis. Transcript wording was used conservatively because room acoustics and automatic speech recognition introduce errors. Where the deck and transcript differed, the deck was treated as the primary formal source and the transcript was used to explain context rather than override it.",
)
add_table(
    doc,
    ["Source", "Contribution", "Status"],
    [
        ("20260918_UnRiskOmega_EN.pdf", "Official company, case, judging, bonus and product material", "Primary source"),
        ("Unriskomega.m4a transcript", "Presentation narrative, sponsor emphasis and spoken context", "Primary conversation source with ASR limitations"),
        ("Unriskomega 2.m4a transcript", "Detailed Q&A on workflow, user segment, sources, briefing design, voice and demo", "Primary conversation source with ASR limitations"),
        ("Unriskomega 3.m4a transcript", "Deployment model, data ownership, Swiss hosting concerns and local-model discussion", "Primary conversation source with ASR limitations"),
        ("START-Hack unriskomega-2026 repository", "Client data, reference data, UI references and ex-custody samples", "Measured technical source at audited commit"),
        ("Prior challenge analysis", "Repository counts, constraints, demo case and acceptance recommendations", "Secondary analysis"),
        ("Prior reverse hackathon playbook", "Build order, pitch, demo reliability and Q&A plan", "Secondary recommendation"),
        ("Prior agent-assist discussion", "Real-time knowledge-assist pattern and evaluation criteria", "Secondary product context"),
    ],
    widths=[2.15, 3.3, 1.25],
    font_size=8.6,
)

add_heading(doc, "Transcript contribution map", 1)
add_table(
    doc,
    ["Recording", "Unique case knowledge retained in this report"],
    [
        ("Unriskomega", "Company narrative, adviser workload, 60-second problem, core and bonus scope, source expectations, judging emphasis and client urgency"),
        ("Unriskomega 2", "Manual trigger requirement, incoming and scheduled use cases, affluent segment, three-part briefing logic, short output preference, voice restrictions, graph toggle, news providers, surprise input and demo behavior"),
        ("Unriskomega 3", "Bank-hosted deployment, separation of product vendor from client data, Swiss control concerns, private or local model options, and the ex-custody integration value"),
    ],
    widths=[1.45, 5.25],
)

add_heading(doc, "Sponsor and reference roles", 1)
add_para(
    doc,
    "The official deck lists Marc Aeberhard as Chief Executive Officer, Michael Nutter CFA as Chief Product Officer, and Frédéric Altorfer as Head Solution Integration and Support. It also presents a summer 2027 internship across product management, business analysis and project management, and application management. Those items are included here for case completeness but are not part of the product requirements.",
)

add_heading(doc, "Final position", 1)
add_para(
    doc,
    "The case is best solved as a trustworthy synthesis product for messy bank data. The team should prove one complete adviser workflow, show the evidence behind the answer, handle unfamiliar input honestly and keep optional features behind the stable core. A concise briefing that exposes freshness and provenance will score more directly against the sponsor's needs than a broader assistant whose claims cannot be audited.",
)

# Keep table rows and headings visually stable.
for table in doc.tables:
    for row in table.rows:
        set_row_cant_split(row)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_together = True
                paragraph.paragraph_format.space_after = Pt(0)

doc.core_properties.title = "UNRISKOMEGA Case Knowledge Base"
doc.core_properties.subject = "From Ping to Pitch in 60 Seconds case report and implementation guide"
doc.core_properties.author = "Benjamin Auer"
doc.core_properties.keywords = "UNRISKOMEGA, URO Advisor Pro, wealth management, AI briefing assistant, START Hack"
doc.save(OUT)
print(OUT)
